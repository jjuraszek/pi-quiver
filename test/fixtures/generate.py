#!/usr/bin/env python3
"""Regenerate the doc_to_md fixtures. Run from the repo root:
uv run --with pymupdf==1.27.2.3 --with openpyxl==3.1.5 --with xlwt --with python-docx --with python-pptx --with pillow --python 3.14 python test/fixtures/generate.py [generator ...]
Pass generator names to regenerate a subset (e.g. charts charts_zero_extent).
Requires soffice on PATH (workbook.xlsx round-trip populates cached formula values)."""
import io, os, shutil, subprocess, tempfile, zipfile, re
import pymupdf
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))

def png_bytes(color):
    buf = io.BytesIO(); Image.new("RGB", (40, 30), color).save(buf, "PNG"); return buf.getvalue()

def multipage_pdf():
    doc = pymupdf.open()
    for n in range(1, 7):
        page = doc.new_page()
        if n == 4:
            continue  # blank page
        page.insert_text((72, 72), f"Chapter {n}" if n % 2 else f"Section {n}", fontsize=20 if n % 2 else 16)
        page.insert_text((72, 110), f"This paragraph belongs to PAGE-{n} and has non-trivial content about wiring.", fontsize=11)
        if n == 2:
            y = 150
            for row in [["Pin", "Color", "Function"], ["1", "Red", "VCC"], ["2", "Black", "GND"]]:
                for i, cell in enumerate(row):
                    page.insert_text((72 + i * 120, y), cell, fontsize=11)
                y += 18
        if n in (3, 5):
            page.insert_image(pymupdf.Rect(72, 200, 272, 350), stream=png_bytes("red" if n == 3 else "blue"))
        if n == 5:  # Add a true inline image (xref 0) alongside the embedded XObject raster.
            raw_image = bytes([0, 255, 0]) * (8 * 6)
            inline_image = b"\nq 60 0 0 45 300 200 cm BI /W 8 /H 6 /CS /RGB /BPC 8 ID\n" + raw_image + b"\nEI Q\n"
            contents = page.get_contents()[0]
            doc.update_stream(contents, doc.xref_stream(contents) + inline_image)
    doc.set_toc([[1, "Chapter 1", 1], [2, "Section 2", 2]])
    doc.set_metadata({"title": "Multipage Fixture", "author": "pi-quiver tests"})
    doc.save(os.path.join(HERE, "multipage.pdf"))
    enc = pymupdf.open(os.path.join(HERE, "multipage.pdf"))
    enc.save(os.path.join(HERE, "encrypted.pdf"), encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="secret", owner_pw="secret")

def shared_resources_pdf():
    doc = pymupdf.open()
    p1 = doc.new_page(); p1.insert_text((72, 72), "PAGE-1 draws the image")
    xref = p1.insert_image(pymupdf.Rect(72, 100, 172, 175), stream=png_bytes("orange"))
    p2 = doc.new_page(); p2.insert_text((72, 72), "PAGE-2 shares the resource but does not draw it")
    # Give page 2 an image resource, then remove its Do operator.
    p2 = doc[1]
    p2.insert_image(pymupdf.Rect(72, 100, 172, 175), stream=png_bytes("orange"))
    for contents in p2.get_contents():
        data = doc.xref_stream(contents)
        doc.update_stream(contents, re.sub(rb"/\S+\s+Do", b"", data))
    doc.save(os.path.join(HERE, "shared-resources.pdf"))

def office():
    from docx import Document
    from docx.enum.text import WD_BREAK
    from pptx import Presentation
    from pptx.util import Inches
    d = Document()
    for n in range(1, 6):
        d.add_heading(f"Heading {n}", level=1)
        d.add_paragraph(f"Body text for PAGE-{n}.")
        if n == 2:
            img = os.path.join(tempfile.gettempdir(), "fx.png"); open(img, "wb").write(png_bytes("purple")); d.add_picture(img)
        if n < 5:
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    d.save(os.path.join(HERE, "multipage.docx"))
    prs = Presentation()
    for n in range(1, 5):
        s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = f"SLIDE-{n}"
        if n == 2:
            img = os.path.join(tempfile.gettempdir(), "fx2.png"); open(img, "wb").write(png_bytes("teal")); s.shapes.add_picture(img, Inches(1), Inches(2))
    prs.save(os.path.join(HERE, "multislide.pptx"))

def workbook():
    import openpyxl, datetime
    from openpyxl.drawing.image import Image as XLImage
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Data"
    ws["A1"] = "Quarterly Data"; ws.merge_cells("A1:C1")
    for r in range(2, 21):
        ws.cell(r, 1, r); ws.cell(r, 2, r * 1.5); ws.cell(r, 3, f"row {r}")
    ws["D17"] = 21; ws["D18"] = "=D17*2"; ws["D19"] = "=SUM(A2:A20)"; ws["E2"] = "pipe|in|text"
    ws["E3"] = datetime.date(2026, 9, 9); ws["E4"] = True; ws["G2"] = "=1/0"
    ws.row_dimensions[4].hidden = True; ws.column_dimensions["F"].hidden = True
    img = os.path.join(tempfile.gettempdir(), "fx3.png"); open(img, "wb").write(png_bytes("gold")); ws.add_image(XLImage(img), "H2")
    st = wb.create_sheet("Settings"); st["C17"] = "threshold"; st["D17"] = 42
    for title in ("A B", "A_B"):
        s = wb.create_sheet(title); s["A1"] = title; s.add_image(XLImage(img), "B2")
    h = wb.create_sheet("Hidden"); h["A1"] = "secret"; h.sheet_state = "hidden"
    raw = os.path.join(tempfile.gettempdir(), "workbook-raw.xlsx"); wb.save(raw)
    out = tempfile.mkdtemp()
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, raw], check=True, capture_output=True, timeout=180)
    conv = os.path.join(out, "workbook-raw.xlsx")
    # strip the cached <v> of D19 so one formula has no cached result
    target = os.path.join(HERE, "workbook.xlsx")
    with zipfile.ZipFile(conv) as zin, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'(<c r="D19"[^>]*>\s*<f[^>]*>[^<]*</f>)\s*<v>[^<]*</v>', rb"\1", data)
            zout.writestr(item, data)

def legacy_xls():
    import xlwt
    wb = xlwt.Workbook(); ws = wb.add_sheet("Legacy")
    ws.write_merge(0, 0, 0, 2, "Legacy Title")
    for r in range(1, 6):
        for c in range(3):
            ws.write(r, c, f"r{r}c{c}")
    ws.write(1, 3, xlwt.Formula("1/0"))
    ws.row(3).hidden = True; ws.col(1).hidden = True
    target = os.path.join(HERE, "legacy.xls")
    with tempfile.TemporaryDirectory() as rawdir, tempfile.TemporaryDirectory() as outdir:
        raw = os.path.join(rawdir, "legacy.xls")
        wb.save(raw)
        subprocess.run(["soffice", "--headless", "--convert-to", "xls", "--outdir", outdir, raw], check=True, capture_output=True, timeout=180)
        shutil.move(os.path.join(outdir, "legacy.xls"), target)

ZERO_EXT = re.compile(rb'<(\w+:)?ext cx="0" cy="0"\s*/>')

def patch_chartsheet_extents(src, dst):
    """openpyxl writes chartsheet drawings with a zero-size absoluteAnchor extent, which LibreOffice honors
    as an empty page. Rewrite the extent at the zip level; never soffice round-trip (LO drops chartsheet drawings)."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if re.fullmatch(r"xl/drawings/drawing\d+\.xml", item.filename):
                data = ZERO_EXT.sub(rb'<\1ext cx="9144000" cy="6858000"/>', data)
            zout.writestr(item, data)

def charts():
    import openpyxl, datetime
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.drawing.image import Image as XLImage
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Data"
    ws.append(["Step", "North", "Double"])
    for r in range(2, 201):
        ws.cell(r, 1, r - 1); ws.cell(r, 2, round((r - 1) * 0.5, 1)); ws.cell(r, 3, f"=A{r}*2")
    line = LineChart(); line.title = "Data trend"; line.add_data(Reference(ws, min_col=2, min_row=1, max_row=200), titles_from_data=True); ws.add_chart(line, "E2")
    img = os.path.join(tempfile.gettempdir(), "fx4.png"); open(img, "wb").write(png_bytes("navy")); ws.add_image(XLImage(img), "E20")
    wb.create_sheet("Empty")
    aux = wb.create_sheet("Aux"); aux.sheet_state = "hidden"
    for r in range(1, 11):
        aux.cell(r, 1, r); aux.cell(r, 2, f"aux {r}")
    cs = wb.create_chartsheet("Trends"); ch = LineChart(); ch.title = "Synthetic trends"
    ch.add_data(Reference(ws, min_col=2, min_row=2, max_row=200)); cs.add_chart(ch)
    cs2 = wb.create_chartsheet("Bars"); bar = BarChart(); bar.title = "Bars"
    bar.add_data(Reference(ws, min_col=1, min_row=2, max_row=200)); cs2.add_chart(bar)
    wide = wb.create_sheet("Wide")
    wide.append([f"C{c}" for c in range(1, 81)])
    words = ["alpha", "beta", "gamma", "delta"]
    for r in range(2, 301):
        row = [r - 1, round((r - 1) / 7, 3), words[r % 4], datetime.date(2026, 1, 1) + datetime.timedelta(days=r)]
        row += [(r * c) % 97 for c in range(5, 81)]
        wide.append(row)
    raw = os.path.join(tempfile.gettempdir(), "charts-raw.xlsx"); wb.save(raw)
    patch_chartsheet_extents(raw, os.path.join(HERE, "charts.xlsx"))

def charts_zero_extent():
    import openpyxl
    from openpyxl.chart import LineChart, Reference
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Data"
    for r in range(1, 4):
        ws.cell(r, 1, r); ws.cell(r, 2, r * 2)
    cs = wb.create_chartsheet("Chart"); ch = LineChart(); ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=3)); cs.add_chart(ch)
    wb.save(os.path.join(HERE, "charts-zero-extent.xlsx"))  # deliberately unpatched: exercises the degenerate-page guard

GENERATORS = {"multipage_pdf": multipage_pdf, "shared_resources_pdf": shared_resources_pdf, "office": office, "workbook": workbook, "legacy_xls": legacy_xls, "charts": charts, "charts_zero_extent": charts_zero_extent}

if __name__ == "__main__":
    import sys
    for name in sys.argv[1:] or GENERATORS:
        GENERATORS[name]()
    print("fixtures written to", HERE)
