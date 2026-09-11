#!/usr/bin/env python3
"""doc_to_md child. argv[1] = mode (info | pdf-primary | pdf-fallback | xlsx | render-pages); options JSON on stdin;
one JSON result on stdout. Exit 0 ok, 1 conversion failure (traceback on stderr), 3 user error
({"error", "pageCount"} on stdout). Library chatter is redirected to stderr so stdout is the result only.
Imports `pymupdf` / `pymupdf4llm` (never the deprecated `fitz` alias)."""
import contextlib
import json
import os
import re
import shutil
import sys
import tempfile
import traceback
import warnings

SEP = "\n\n--- end of page.page_number={n} ---\n\n"
DEGRADED_NOTE = "degraded: PyMuPDF text extraction - layout/tables not preserved"
MARKDOWN_IMAGE_RE = re.compile(r"(!\[[^\]]*\]\(\s*)(?:<([^>]+)>|([^)]*?))(\s*\))")
HTML_IMAGE_RE = re.compile(
    r"(<img\b[^>]*?\bsrc\s*=\s*)(?:\"([^\"]*)\"|'([^']*)'|([^\s\"'=<>`]+))", re.IGNORECASE
)


def image_source_map(source, target, filename):
    real = os.path.realpath(source)
    return {source: target, source.replace("\\", "/"): target,
            real: target, real.replace("\\", "/"): target, filename: target}


def rewrite_image_destinations(md, sources):
    def markdown(match):
        dest = match.group(2) if match.group(2) is not None else match.group(3)
        target = sources.get(dest)
        if target is None:
            return match.group(0)
        if match.group(2) is not None:
            return f"{match.group(1)}<{target}>{match.group(4)}"
        return f"{match.group(1)}{target}{match.group(4)}"

    def html(match):
        dest = next(value for value in match.groups()[1:] if value is not None)
        target = sources.get(dest)
        if target is None:
            return match.group(0)
        quote = '"' if match.group(2) is not None else "'" if match.group(3) is not None else ""
        return f"{match.group(1)}{quote}{target}{quote}"

    return HTML_IMAGE_RE.sub(html, MARKDOWN_IMAGE_RE.sub(markdown, md))


def user_error(msg, page_count=None):
    json.dump({"error": msg, **({"pageCount": page_count} if page_count is not None else {})}, sys.stdout)
    return 3


def check_pages(pages, page_count):
    if pages is None:
        return list(range(1, page_count + 1))
    bad = [p for p in pages if p < 1 or p > page_count]
    if bad:
        raise UserError(f"pages out of range: {', '.join(map(str, bad))} (document has {page_count} pages)", page_count)
    return pages


class UserError(Exception):
    def __init__(self, msg, page_count=None):
        super().__init__(msg)
        self.page_count = page_count


def open_pdf(path):
    import pymupdf
    doc = pymupdf.open(path)
    if doc.needs_pass:
        raise UserError("Password-protected PDF", doc.page_count)
    return doc


def page_dir(staging, n):
    d = os.path.join(staging, f"p{n}")
    os.makedirs(d, exist_ok=True)
    return d


def mark_done(d):
    open(os.path.join(d, ".done"), "w").close()


def mode_info(o):
    import pymupdf  # noqa: F401
    doc = open_pdf(o["path"])
    meta = {k: v for k, v in (doc.metadata or {}).items() if v}
    toc = [[lvl, title, page] for lvl, title, page in doc.get_toc()]
    return {"pageCount": doc.page_count, "metadata": meta, "toc": toc}


def mode_pdf_primary(o):
    import pymupdf4llm
    doc = open_pdf(o["path"])
    pages = check_pages(o.get("pages"), doc.page_count)
    staging, out, empty, failed, notes = o["stagingDir"], [], [], [], []
    for n in pages:
        d = page_dir(staging, n)
        try:
            # Space-free temp dir: pymupdf4llm's md_path() mangles paths containing spaces/parens.
            with tempfile.TemporaryDirectory() as tmp:
                md = pymupdf4llm.to_markdown(doc, pages=[n - 1], write_images=True, image_path=tmp,
                                               image_format=o["imageFormat"], dpi=o["imageDpi"],
                                               use_ocr=False, page_separators=False)
                sources = {}
                for i, f in enumerate(sorted(os.listdir(tmp)), 1):
                    dest = f"img{i}{os.path.splitext(f)[1].lower()}"
                    source = os.path.join(tmp, f)
                    target = f"p{n}/{dest}"
                    sources.update(image_source_map(source, target, f))
                    os.replace(source, os.path.join(d, dest))
                md = rewrite_image_destinations(md, sources)
            if not md.strip():
                empty.append(n)
            out.append(md.rstrip())
            mark_done(d)
        except Exception as exc:  # noqa: BLE001
            shutil.rmtree(d, ignore_errors=True)
            failed.append({"page": n, "error": f"{type(exc).__name__}: {exc}"[:300]})
            empty.append(n)
            out.append("")
        out.append(SEP.format(n=n).strip("\n"))
    if failed and len(failed) == len(pages):
        raise RuntimeError("every selected page failed: " + failed[0]["error"])
    return {"markdown": "\n\n".join(out) + "\n", "pages": pages, "pageCount": doc.page_count,
            "emptyPages": empty, "failedPages": failed, "notes": notes}


def mode_pdf_fallback(o):
    import pymupdf
    doc = open_pdf(o["path"])
    pages = check_pages(o.get("pages"), doc.page_count)
    keep = {int(k): v for k, v in (o.get("keepPages") or {}).items()}
    staging, out, empty, failed = o["stagingDir"], [], [], []
    for n in pages:
        links = [f"![](images/{f})" for f in keep.get(n, [])]
        text = ""
        try:
            page = doc[n - 1]
            text = page.get_text("text").strip()
            if n not in keep:
                d = page_dir(staging, n)
                i = 0
                for info in page.get_image_info(xrefs=True):
                    i += 1
                    xref = info.get("xref", 0)
                    if xref > 0:
                        img = doc.extract_image(xref)
                        name = f"img{i}.{img['ext'].lower()}"
                        with open(os.path.join(d, name), "wb") as fh:
                            fh.write(img["image"])
                    else:
                        name = f"img{i}.{o['imageFormat']}"
                        page.get_pixmap(clip=pymupdf.Rect(info["bbox"]), dpi=o["imageDpi"]).save(os.path.join(d, name))
                    links.append(f"![](p{n}/{name})")
                mark_done(d)
        except Exception as exc:  # noqa: BLE001
            shutil.rmtree(os.path.join(staging, f"p{n}"), ignore_errors=True)
            text = ""
            links = [f"![](images/{f})" for f in keep.get(n, [])]
            failed.append({"page": n, "error": f"{type(exc).__name__}: {exc}"[:300]})
        if not text:
            empty.append(n)
        out.append("\n\n".join(x for x in [text, "\n".join(links)] if x))
        out.append(SEP.format(n=n).strip("\n"))
    if failed and len(failed) == len(pages):
        raise RuntimeError("every selected page failed: " + failed[0]["error"])
    return {"markdown": "\n\n".join(out) + "\n", "pages": pages, "pageCount": doc.page_count,
            "emptyPages": empty, "failedPages": failed, "notes": [DEGRADED_NOTE]}


def esc(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        return repr(v)
    import datetime
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.isoformat()
    return str(v).replace("\\", "\\\\").replace("|", "\\|").replace("\r\n", "<br>").replace("\n", "<br>")


def col_letter(i):
    from openpyxl.utils import get_column_letter
    return get_column_letter(i)


PREVIEW_ROWS, PREVIEW_COLS = 100, 50
PROFILE_MAJORITY, DISTINCT_CAP, SLUG_MAX = 0.6, 50, 40
MAX_RENDER_PX, MIN_RENDER_DPI, MIN_PAGE_PT = 16_000_000, 36, 72
INV_HEADER = "| # | name | kind | size | hidden | charts | images | rendered | data |\n|---|---|---|---|---|---|---|---|---|\n"
XLS_NOTE = "Rendered views: unavailable (visual detection not supported for .xls)"


def slug(name):
    s = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")[:SLUG_MAX].strip("-")
    return s or "sheet"


def csv_value(v):
    import datetime
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.isoformat()
    return str(v)


def cell_class(v):
    import datetime
    if isinstance(v, bool):
        return "bool"
    if isinstance(v, int):
        return "int"
    if isinstance(v, float):
        return "float"
    if isinstance(v, (datetime.date, datetime.datetime)):
        return "date"
    return "str"


def is_formula(src):
    return isinstance(src, str) and src.startswith("=")


class ColProfile:
    def __init__(self):
        self.n, self.classes, self.lo, self.hi, self.distinct = 0, {}, None, None, set()

    def add(self, cls, v):
        self.n += 1
        self.classes[cls] = self.classes.get(cls, 0) + 1
        if cls in ("int", "float", "date"):
            try:
                if self.lo is None or v < self.lo:
                    self.lo = v
                if self.hi is None or v > self.hi:
                    self.hi = v
            except TypeError:
                pass
        if len(self.distinct) <= DISTINCT_CAP:
            self.distinct.add(repr(v))

    def row(self, letter, header):
        if not self.n:
            return f"| {letter} | {header} | - | 0 | - | - | 0 |"
        cls, cnt = max(self.classes.items(), key=lambda kv: kv[1])
        typ = cls if cnt / self.n >= PROFILE_MAJORITY else "mixed"
        rng = typ in ("int", "float", "date") and self.lo is not None
        lo, hi = (esc(self.lo), esc(self.hi)) if rng else ("-", "-")
        distinct = f">{DISTINCT_CAP}" if len(self.distinct) > DISTINCT_CAP else str(len(self.distinct))
        return f"| {letter} | {header} | {typ} | {self.n} | {lo} | {hi} | {distinct} |"


def chart_line(ch):
    title = "untitled"
    tx = getattr(getattr(ch, "title", None), "tx", None)
    if tx is not None:
        if getattr(tx, "rich", None) is not None:
            title = "".join((r.t or "") for p in tx.rich.p for r in (p.r or [])) or "untitled"
        elif getattr(tx, "strRef", None) is not None:
            title = tx.strRef.f or "untitled"
    series = list(getattr(ch, "series", None) or [])
    refs = [s.val.numRef.f for s in series if getattr(s, "val", None) is not None and s.val.numRef is not None]
    shown = ", ".join(refs[:5]) + (", ..." if len(refs) > 5 else "")
    return f'- {type(ch).__name__} "{esc(title)}" - {len(series)} series ({shown})'


def write_csv(idx, title, cells, R, C, csv_dir):
    if not R:
        return None
    import csv
    os.makedirs(csv_dir, exist_ok=True)
    name = f"s{idx}-{slug(title)}.csv"
    with open(os.path.join(csv_dir, name), "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        for row in cells[:R]:
            w.writerow([csv_value(val if val is not None else src) for src, val in row[:C]])
    return f"sheets/{name}"


def data_line(R, C, csv_rel, formulas):
    if not R:
        return "Data: none"
    return f"Data: [{csv_rel}]({csv_rel}) - {R} rows x {C} cols" + (f", {formulas} formulas" if formulas else "")


def preview_and_profile(lines, cells, profiles, R, C, continuation):
    r_lim, c_lim = min(R, PREVIEW_ROWS), min(C, PREVIEW_COLS)
    truncated = r_lim < R or c_lim < C
    lines.append("")
    lines.append(f"Preview (rows 1-{r_lim} of {R}, cols A-{col_letter(c_lim)} of {C}) - full data in the CSV above:" if truncated else f"Content ({R} rows x {C} cols):")
    lines += ["| | " + " | ".join(col_letter(c) for c in range(1, c_lim + 1)) + " |", "|---|" + "---|" * c_lim]
    for r in range(1, r_lim + 1):
        out = []
        for c in range(1, c_lim + 1):
            if (r, c) in continuation:
                out.append("")
                continue
            src, val = cells[r - 1][c - 1]
            if is_formula(src):
                out.append(f"{esc(val) if val is not None else '(no cached result)'} ({esc(src)})")
            else:
                out.append(esc(src))
        lines.append(f"| {r} | " + " | ".join(out) + " |")
    if not truncated:
        return
    lines += ["", "Columns:", "| col | header | type | non-empty | min | max | distinct |", "|---|---|---|---|---|---|---|"]
    for c, p in enumerate(profiles[:C], 1):
        h = cells[0][c - 1][0]
        lines.append(p.row(col_letter(c), esc(h) if isinstance(h, str) and not is_formula(h) else "-"))


def inv_row(idx, name, kind, size, hidden, charts, images, rvs, data):
    return f"| {idx} | {esc(name)} | {kind} | {size} | {'yes' if hidden else 'no'} | {charts} | {images} | {rvs} | {data} |"


def sheet_info(idx, name, kind, hidden, R, C, hr, hc, charts, images, csv_rel):
    return {"index": idx, "name": name, "kind": kind, "hidden": hidden, "rows": R, "cols": C, "hiddenRows": hr, "hiddenCols": hc,
            "charts": charts, "images": images, "rendered": False, "csv": csv_rel}

def mode_xlsx(o):
    path, staging, csv_dir = o["path"], o["stagingDir"], o["sheetsStagingDir"]
    if path.lower().endswith(".xls"):
        return mode_xls(o)
    import openpyxl
    from openpyxl.chartsheet import Chartsheet
    wb_f = openpyxl.load_workbook(path, data_only=False)
    wb_v = openpyxl.load_workbook(path, data_only=True)
    stem = os.path.splitext(os.path.basename(path))[0]
    notes, images, sheets, render, inv, sections = [], [], [], [], [], []
    for idx, ws in enumerate(wb_f._sheets):
        hidden = ws.sheet_state != "visible"
        charts = list(getattr(ws, "_charts", []) or [])
        if isinstance(ws, Chartsheet):
            lines = [f"## {esc(ws.title)} (chartsheet)"]
            if charts:
                lines += ["Charts:"] + [chart_line(ch) for ch in charts]
                render.append(idx)
                lines.append(f"<!--rv:{idx}-->")
            inv.append(inv_row(idx, ws.title, "chartsheet", "-", hidden, len(charts), 0, f"<!--rvs:{idx}-->" if charts else "-", "-"))
            sheets.append(sheet_info(idx, ws.title, "chartsheet", hidden, None, None, 0, 0, len(charts), 0, None))
            sections.append("\n".join(lines))
            continue
        wv = wb_v[ws.title]
        raw_r, raw_c = ws.max_row, ws.max_column
        cells, formulas, R, C = [], 0, 0, 0
        profiles = [ColProfile() for _ in range(raw_c)]
        for r, (fr, vr) in enumerate(zip(ws.iter_rows(min_row=1, max_row=raw_r, max_col=raw_c),
                                         wv.iter_rows(min_row=1, max_row=raw_r, max_col=raw_c)), 1):
            row = []
            for c, (fc, vc) in enumerate(zip(fr, vr), 1):
                src, val = fc.value, vc.value
                if src is not None:
                    R, C = r, max(C, c)
                    if is_formula(src):
                        formulas += 1
                        profiles[c - 1].add("formula" if val is None else cell_class(val), val)
                    else:
                        profiles[c - 1].add(cell_class(src), src)
                row.append((src, val))
            cells.append(row)
        lines = [f"## {esc(ws.title)}"]
        if hidden:
            lines.append("Hidden sheet")
        merged = [str(r) for r in ws.merged_cells.ranges]
        if merged:
            lines.append("Merged: " + ", ".join(merged))
        hr = [str(r) for r, dim in ws.row_dimensions.items() if dim.hidden]
        if hr:
            lines.append("Hidden rows: " + ",".join(hr))
        hc = []
        for dim in ws.column_dimensions.values():
            if dim.hidden:
                hc += [col_letter(i) for i in range(dim.min, dim.max + 1)]
        if hc:
            lines.append("Hidden cols: " + ",".join(hc))
        csv_rel = write_csv(idx, ws.title, cells, R, C, csv_dir)
        lines.append(data_line(R, C, csv_rel, formulas))
        if charts:
            lines += ["Charts:"] + [chart_line(ch) for ch in charts]
        raw_images = list(getattr(ws, "_images", []) or [])
        img_lines = []
        for i, img in enumerate(raw_images, 1):
            try:
                fmt = (getattr(img, "format", None) or "png").lower()
                name = f"s{idx}-{i}.{fmt}"
                data = img._data() if callable(getattr(img, "_data", None)) else img.ref.getvalue()
                with open(os.path.join(staging, name), "wb") as fh:
                    fh.write(data)
                images.append({"sheetIndex": idx, "file": name})
                img_lines.append(f"- ![{esc(ws.title)} image {i}]({name})")
            except Exception as exc:  # noqa: BLE001
                notes.append(f"sheet {ws.title}: image {i} not extracted ({type(exc).__name__})")
        if img_lines:
            lines += ["Images:"] + img_lines
        visual = bool(charts or raw_images)
        if visual:
            render.append(idx)
            lines.append(f"<!--rv:{idx}-->")
        if R:
            continuation = set()
            for rng in ws.merged_cells.ranges:
                for r in range(rng.min_row, rng.max_row + 1):
                    for c in range(rng.min_col, rng.max_col + 1):
                        if (r, c) != (rng.min_row, rng.min_col):
                            continuation.add((r, c))
            preview_and_profile(lines, cells, profiles, R, C, continuation)
        inv.append(inv_row(idx, ws.title, "worksheet", f"{R} x {C}", hidden, len(charts), len(raw_images),
                           f"<!--rvs:{idx}-->" if visual else "-", f"[{csv_rel}]({csv_rel})" if csv_rel else "-"))
        sheets.append(sheet_info(idx, ws.title, "worksheet", hidden, R, C, len(hr), len(hc), len(charts), len(raw_images), csv_rel))
        sections.append("\n".join(lines))
    md = f"# {esc(stem)}\n\n## Sheets\n" + INV_HEADER + "\n".join(inv) + "\n\n" + "\n\n".join(sections) + "\n"
    return {"markdown": md, "images": images, "notes": notes, "sheets": sheets, "renderPages": render, "sheetCount": len(wb_f._sheets)}

def mode_xls(o):
    import xlrd
    path, csv_dir = o["path"], o["sheetsStagingDir"]
    book = xlrd.open_workbook(path, formatting_info=True, logfile=sys.stderr)
    stem = os.path.splitext(os.path.basename(path))[0]
    sheets, inv, sections = [], [], []
    for idx in range(book.nsheets):
        sh = book.sheet_by_index(idx)
        hidden = sh.visibility != 0
        cells, R, C = [], 0, 0
        profiles = [ColProfile() for _ in range(sh.ncols)]
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                v = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    v = xlrd.xldate_as_datetime(v, book.datemode)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    v = bool(v)
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    v = xlrd.error_text_from_code.get(int(v), f"#ERR{int(v)}")
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    v = None
                if v is not None:
                    R, C = r + 1, max(C, c + 1)
                    profiles[c].add(cell_class(v), v)
                row.append((v, v))
            cells.append(row)
        lines = [f"## {esc(sh.name)}", "Formulas: unavailable (.xls via xlrd); Images: unavailable"]
        if hidden:
            lines.append("Hidden sheet")
        merged = [f"{col_letter(c0 + 1)}{r0 + 1}:{col_letter(c1)}{r1}" for r0, r1, c0, c1 in sh.merged_cells]
        if merged:
            lines.append("Merged: " + ", ".join(merged))
        hr = [str(r + 1) for r, info in sh.rowinfo_map.items() if info.hidden]
        if hr:
            lines.append("Hidden rows: " + ",".join(hr))
        hc = [col_letter(c + 1) for c, info in sh.colinfo_map.items() if info.hidden]
        if hc:
            lines.append("Hidden cols: " + ",".join(hc))
        csv_rel = write_csv(idx, sh.name, cells, R, C, csv_dir)
        lines.append(data_line(R, C, csv_rel, 0))
        if R:
            continuation = {(r + 1, c + 1) for r0, r1, c0, c1 in sh.merged_cells for r in range(r0, r1) for c in range(c0, c1) if (r, c) != (r0, c0)}
            preview_and_profile(lines, cells, profiles, R, C, continuation)
        inv.append(inv_row(idx, sh.name, "worksheet", f"{R} x {C}", hidden, 0, 0, "-", f"[{csv_rel}]({csv_rel})" if csv_rel else "-"))
        sheets.append(sheet_info(idx, sh.name, "worksheet", hidden, R, C, len(hr), len(hc), 0, 0, csv_rel))
        sections.append("\n".join(lines))
    md = f"# {esc(stem)}\n\n## Sheets\n" + INV_HEADER + "\n".join(inv) + "\n\n" + XLS_NOTE + "\n\n" + "\n\n".join(sections) + "\n"
    return {"markdown": md, "images": [], "notes": [XLS_NOTE], "sheets": sheets, "renderPages": [], "sheetCount": book.nsheets}

def mode_info_excel(o):
    path = o["path"]
    sheets = []
    if path.lower().endswith(".xls"):
        import xlrd
        book = xlrd.open_workbook(path, formatting_info=True, logfile=sys.stderr)
        for idx in range(book.nsheets):
            sh = book.sheet_by_index(idx)
            sheets.append(sheet_info(idx, sh.name, "worksheet", sh.visibility != 0, sh.nrows, sh.ncols,
                                     sum(1 for i in sh.rowinfo_map.values() if i.hidden), sum(1 for i in sh.colinfo_map.values() if i.hidden), 0, 0, None))
    else:
        import openpyxl
        from openpyxl.chartsheet import Chartsheet
        wb = openpyxl.load_workbook(path, data_only=True)
        for idx, ws in enumerate(wb._sheets):
            hidden = ws.sheet_state != "visible"
            charts = len(getattr(ws, "_charts", []) or [])
            if isinstance(ws, Chartsheet):
                sheets.append(sheet_info(idx, ws.title, "chartsheet", hidden, None, None, 0, 0, charts, 0, None))
                continue
            hc = sum(d.max - d.min + 1 for d in ws.column_dimensions.values() if d.hidden)
            hr = sum(1 for d in ws.row_dimensions.values() if d.hidden)
            sheets.append(sheet_info(idx, ws.title, "worksheet", hidden, ws.max_row, ws.max_column, hr, hc, charts, len(getattr(ws, "_images", []) or []), None))
    return {"sheets": sheets}

def mode_render_pages(o):
    import math
    import pymupdf
    doc = pymupdf.open(o["path"])
    expected = o["expectedPages"]
    if doc.page_count != expected:
        return {"ok": False, "reason": f"page-count mismatch ({doc.page_count} vs {expected})"}
    fmt, dpi, staging = o["imageFormat"], o["imageDpi"], o["stagingDir"]
    rendered, failed = [], []
    for idx in o["sheetIndices"]:
        try:
            page = doc[idx]
            w, h = page.rect.width, page.rect.height
            if w < MIN_PAGE_PT or h < MIN_PAGE_PT:
                failed.append({"idx": idx, "reason": f"rendered view degenerate (page {w:.0f} x {h:.0f} pt)"})
                continue
            eff = min(dpi, math.floor(math.sqrt(MAX_RENDER_PX / (w * h / 72 ** 2))))
            if eff < MIN_RENDER_DPI:
                failed.append({"idx": idx, "reason": f"rendered view too large (page {w:.0f} x {h:.0f} pt)"})
                continue
            name = f"s{idx}.{fmt}"
            page.get_pixmap(dpi=eff).save(os.path.join(staging, name))
            rendered.append({"idx": idx, "file": name, "dpi": eff})
        except Exception as exc:  # noqa: BLE001
            failed.append({"idx": idx, "reason": f"render failed: {type(exc).__name__}: {exc}"[:300]})
    return {"ok": True, "rendered": rendered, "failed": failed}


MODES = {"info": None, "pdf-primary": mode_pdf_primary, "pdf-fallback": mode_pdf_fallback, "xlsx": mode_xlsx, "render-pages": mode_render_pages}


def main():
    if len(sys.argv) != 2 or sys.argv[1] not in MODES:
        print("usage: doc_to_md.py <info|pdf-primary|pdf-fallback|xlsx|render-pages>  (options JSON on stdin)", file=sys.stderr)
        return 1
    mode = sys.argv[1]
    o = json.loads(sys.stdin.read() or "{}")
    try:
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            with contextlib.redirect_stdout(sys.stderr):
                if mode == "info":
                    result = mode_info_excel(o) if o["path"].lower().endswith((".xlsx", ".xls")) else mode_info(o)
                else:
                    result = MODES[mode](o)
        if "notes" in result:
            result["notes"] += sorted({str(w.message) for w in caught})
    except UserError as ue:
        return user_error(str(ue), ue.page_count)
    except Exception:  # noqa: BLE001
        traceback.print_exc()
        return 1
    json.dump(result, sys.stdout)
    return 0


if __name__ == "__main__":
    sys.exit(main())
